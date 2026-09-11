"""Раздел товаров: числа, отсев помеченных отгрузок и цена расчёта.

Агрегаты здесь считаются условием прямо в `Sum`/`Count`, а не подзапросом
`shipmentitem__in=...`, как раньше. У замены есть цена: подзапрос строился
от `ShipmentItem.objects`, и отсев отгрузок, пропавших из МойСклад, приезжал
вместе с менеджером. Обход связи менеджера не спрашивает, поэтому условие
выписано руками — и должно остаться выписанным.
"""

import json
from datetime import timedelta
from io import BytesIO
from decimal import Decimal

from django.contrib.auth.models import User
from django.db import connection
from django.test import Client, TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from core.models import (
    Counterparty, Product, SalesChannel, Shipment, ShipmentItem,
)


class ProductListTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        user = User.objects.create_user(username='products', password='pass12345')
        from api.access import grant_all_assignable_pages
        grant_all_assignable_pages(user)
        self.client.login(username='products', password='pass12345')

        self.now = timezone.now()
        self.customer = Counterparty.objects.create(name='Покупатель', external_id='c-1')
        self.channel = SalesChannel.objects.create(name='Сайт', external_id='ch-1')

    def ship(self, product, quantity, price, days_ago=1, channel=None, external_id=None):
        shipment = Shipment.objects.create(
            counterparty=self.customer,
            date=self.now - timedelta(days=days_ago),
            external_id=external_id or f'sh-{Shipment.all_objects.count()}',
            sales_channel=channel,
        )
        ShipmentItem.objects.create(
            shipment=shipment, product=product,
            quantity=Decimal(quantity), price=Decimal(price),
        )
        return shipment

    def data(self, query=''):
        response = self.client.get(f'/api/products/{query}')
        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertEqual(payload['status'], 'success')
        return payload['data']


class ProductAggregateTests(ProductListTestCase):
    def setUp(self):
        super().setUp()
        self.product = Product.objects.create(
            name='Вакцина', external_id='p-1', group='Товары', subgroup='Биопрепараты',
        )
        self.ship(self.product, '10', '100', days_ago=1)
        self.ship(self.product, '5', '200', days_ago=9)

    def test_totals_are_summed_over_shipments(self):
        row = self.data()['products'][0]
        self.assertEqual(row['quantity'], 15)
        self.assertEqual(row['total_sum'], 10 * 100 + 5 * 200)
        self.assertEqual(row['average_price'], 150)
        self.assertEqual(row['shipments_count'], 2)

    def test_shipment_marked_as_missing_is_left_out(self):
        """Отгрузка, пропавшая из МойСклад, перестаёт считаться."""
        missing = Shipment.all_objects.order_by('pk').first()
        missing.deleted_at = timezone.now()
        missing.save(update_fields=['deleted_at'])

        row = self.data()['products'][0]
        self.assertEqual(row['quantity'], 5)
        self.assertEqual(row['total_sum'], 1000)
        self.assertEqual(row['shipments_count'], 1)

    def test_period_narrows_the_totals(self):
        recent = (self.now - timedelta(days=3)).date().isoformat()
        row = self.data(f'?startDate={recent}')['products'][0]
        self.assertEqual(row['quantity'], 10)
        self.assertEqual(row['shipments_count'], 1)

    def test_sales_channel_narrows_the_totals(self):
        other = Product.objects.create(
            name='Сыворотка', external_id='p-2', group='Товары', subgroup='Биопрепараты',
        )
        self.ship(other, '3', '50', days_ago=2, channel=self.channel)

        data = self.data('?salesChannel=Сайт')
        self.assertEqual([row['name'] for row in data['products']], ['Сыворотка'])
        self.assertEqual(data['products'][0]['quantity'], 3)

    def test_stats_top_lists_match_the_rows(self):
        other = Product.objects.create(
            name='Сыворотка', external_id='p-2', group='Товары', subgroup='Биопрепараты',
        )
        self.ship(other, '1', '10', days_ago=2)

        stats = self.data()['stats']
        self.assertEqual(stats['total_products'], 2)
        self.assertEqual(stats['top_by_quantity'][0]['name'], 'Вакцина')
        self.assertEqual(stats['top_by_revenue'][0]['name'], 'Вакцина')
        self.assertEqual(stats['top_by_average_quantity'][0]['name'], 'Вакцина')

    def test_product_without_shipments_shows_zeroes(self):
        Product.objects.create(
            name='Новинка', external_id='p-3', group='Товары', subgroup='Биопрепараты',
        )
        rows = {row['name']: row for row in self.data()['products']}
        self.assertEqual(rows['Новинка']['quantity'], 0)
        self.assertEqual(rows['Новинка']['total_sum'], 0)
        self.assertEqual(rows['Новинка']['shipments_count'], 0)


class ExportMatchesScreenTests(ProductListTestCase):
    """Выгрузка в Excel обязана показывать то же, что экран раздела.

    Расчёт у них общий (`aggregated_products`), и это не украшение: пока
    у выгрузки была своя копия агрегатов, любое новое условие отбора
    приходилось дописывать дважды, а разъехавшись, они дали бы разные числа
    за один и тот же период.
    """

    def setUp(self):
        super().setUp()
        self.recent = Product.objects.create(
            name='Свежий', external_id='p-1', group='Товары', subgroup='Биопрепараты',
        )
        self.old = Product.objects.create(
            name='Давний', external_id='p-2', group='Товары', subgroup='Биопрепараты',
        )
        self.ship(self.recent, '10', '100', days_ago=1)
        self.ship(self.old, '5', '200', days_ago=90)

    def exported_names(self, query=''):
        response = self.client.get(f'/api/products/export/{query}')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        names = []
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                names.append(row[0])
        return sorted(names)

    def test_export_without_filters_matches_the_screen(self):
        screen = sorted(row['name'] for row in self.data()['products'])
        self.assertEqual(self.exported_names(), screen)

    def test_export_with_a_period_matches_the_screen(self):
        recent = (self.now - timedelta(days=30)).date().isoformat()
        screen = sorted(row['name'] for row in self.data(f'?startDate={recent}')['products'])
        self.assertEqual(screen, ['Свежий'])
        self.assertEqual(self.exported_names(f'?startDate={recent}'), screen)


class ProductPagingTests(ProductListTestCase):
    def test_paging_does_not_repeat_or_lose_rows(self):
        """Товары без отгрузок неразличимы по сумме — порядок обязан быть устойчивым."""
        for index in range(4):
            Product.objects.create(
                name=f'Товар {index}', external_id=f'p-{index}',
                group='Товары', subgroup='Биопрепараты',
            )
        first = [row['name'] for row in self.data('?pageSize=2&page=1')['products']]
        second = [row['name'] for row in self.data('?pageSize=2&page=2')['products']]
        self.assertEqual(len(set(first + second)), 4)


class ProductQueryCountTests(ProductListTestCase):
    def create_products(self, count):
        Product.objects.all().delete()
        for index in range(count):
            product = Product.objects.create(
                name=f'Товар {index}', external_id=f'bulk-{index}',
                group='Товары', subgroup='Биопрепараты',
            )
            self.ship(product, '1', '10', days_ago=1)

    def queries_for(self, product_count):
        self.create_products(product_count)
        with CaptureQueriesContext(connection) as captured:
            self.data('?pageSize=10')
        return len(captured)

    def test_query_count_does_not_grow_with_the_catalogue(self):
        for_five = self.queries_for(5)
        for_fifty = self.queries_for(50)
        self.assertEqual(
            for_five, for_fifty,
            f'запросов на 5 товаров: {for_five}, на 50: {for_fifty}',
        )


class FilterDictionaryTests(ProductListTestCase):
    """Справочники панели фильтров отдаются отдельно от расчёта."""

    def test_product_filters_return_subgroups_and_channels(self):
        product = Product.objects.create(
            name='Вакцина', external_id='p-1', group='Товары', subgroup='Биопрепараты',
        )
        self.ship(product, '1', '10', channel=self.channel)

        response = self.client.get('/api/products/filters/')
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)['data']
        self.assertEqual(data['available_subgroups'], ['Биопрепараты'])
        self.assertEqual(data['available_sales_channels'], ['Сайт'])

    def test_product_filters_do_not_aggregate(self):
        """Справочник не считает отгрузки: два списка строк — два запроса к данным.

        Точное число запросов задаёт не вью, а сессия и проверка прав, поэтому
        сравниваем не с константой, а сами с собой на разных объёмах.
        """
        def queries_for(product_count):
            Product.objects.all().delete()
            for index in range(product_count):
                product = Product.objects.create(
                    name=f'Товар {index}', external_id=f'p-{index}',
                    group='Товары', subgroup=f'Подгруппа {index % 3}',
                )
                self.ship(product, '1', '10')
            with CaptureQueriesContext(connection) as captured:
                self.client.get('/api/products/filters/')
            return len(captured)

        self.assertEqual(queries_for(3), queries_for(30))

    def test_supply_material_filters_return_groups(self):
        """Панель фильтров поставок берёт группы отдельно от пересчёта поставок."""
        response = self.client.get('/api/supplies/materials/filters/')
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)['data']
        self.assertIn('Тара', data['available_groups'])
        self.assertNotIn('materials', data)

    def test_material_filters_return_groups_and_suppliers(self):
        response = self.client.get('/api/materials/filters/')
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)['data']
        self.assertIn('Тара', data['available_groups'])
        self.assertEqual(data['counterparties'], [])
